"""Test decks made with python-pptx (an independent writer of .pptx)."""
import sys
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

out = sys.argv[1]
p = Presentation()
p.slide_width, p.slide_height = Emu(12192000), Emu(6858000)
s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = "Quarterly Review"
s.placeholders[1].text = "Folio Deliver test deck"
s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = "Agenda"
tf = s.placeholders[1].text_frame
tf.text = "Results"
for t, lvl in [("Revenue up 12%", 1), ("Costs flat", 1), ("Plans", 0), ("Hiring", 1), ("Questions", 0)]:
    para = tf.add_paragraph(); para.text = t; para.level = lvl
s = p.slides.add_slide(p.slide_layouts[5])
s.shapes.title.text = "Shapes and a table"
sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.8), Inches(3), Inches(1.5))
sh.text = "Rounded"
sh2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(1.8), Inches(2), Inches(2))
sh2.fill.solid(); sh2.fill.fore_color.rgb = RGBColor(0xFF, 0x82, 0x00)
sh2.rotation = 20
tb = s.shapes.add_textbox(Inches(0.5), Inches(4), Inches(5), Inches(1))
r = tb.text_frame.paragraphs[0].add_run(); r.text = "Bold red 28pt"; r.font.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor(0xC0, 0, 0)
r2 = tb.text_frame.paragraphs[0].add_run(); r2.text = " and italic"; r2.font.italic = True; r2.font.size = Pt(28)
t = s.shapes.add_table(3, 3, Inches(7), Inches(1.8), Inches(5.5), Inches(1.5)).table
for i in range(3):
    for j in range(3):
        t.cell(i, j).text = "H%d" % j if i == 0 else "%d" % (i * j)
s = p.slides.add_slide(p.slide_layouts[5])
s.shapes.title.text = "A chart"
cd = CategoryChartData(); cd.categories = ["Q1", "Q2", "Q3", "Q4"]; cd.add_series("Sales", (3, 5, 4, 7))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(8), Inches(5), cd)
p.save(out)
